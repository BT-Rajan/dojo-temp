"""
PDF and Excel exports (decision P2-16). Every report is turned into a simple "document": a title,
a subtitle and a list of blocks — ("table", header, rows), ("text", paragraph) or ("heading", text) —
and written to either format. Arabic is shaped and laid out right-to-left in the PDF.
"""

import io
from decimal import Decimal
from pathlib import Path

from django.utils import timezone

FONT_DIR = Path(__file__).resolve().parent / "fonts"
_fonts_ready = False


def _register_fonts():
    global _fonts_ready
    if _fonts_ready:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    pdfmetrics.registerFont(TTFont("Plex", str(FONT_DIR / "IBMPlexSansArabic-Regular.ttf")))
    pdfmetrics.registerFont(TTFont("Plex-Bold", str(FONT_DIR / "IBMPlexSansArabic-Bold.ttf")))
    _fonts_ready = True


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return f"{value:.3f}"
    if isinstance(value, bool):
        return "✓" if value else "—"
    if hasattr(value, "strftime"):
        return value.strftime("%d/%m/%Y")
    return str(value)


def _shape(text: str, rtl: bool) -> str:
    """Arabic letters joined and ordered for display (the PDF library draws left to right)."""
    import arabic_reshaper
    from bidi.algorithm import get_display

    if not text:
        return ""
    if any("؀" <= ch <= "ۿ" for ch in text):
        return get_display(arabic_reshaper.reshape(text), base_dir="R" if rtl else "L")
    return text


def to_pdf(doc: dict, lang="en") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    _register_fonts()
    rtl = lang == "ar"
    align = TA_RIGHT if rtl else TA_LEFT
    wide = any(b[0] == "table" and len(b[1]) > 7 for b in doc["blocks"])
    pagesize = landscape(A4) if wide else A4
    buf = io.BytesIO()
    pdf = SimpleDocTemplate(
        buf,
        pagesize=pagesize,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=doc["title"],
        author="Elite Academy",
    )
    title = ParagraphStyle("t", fontName="Plex-Bold", fontSize=15, leading=20, alignment=align)
    sub = ParagraphStyle(
        "s", fontName="Plex", fontSize=9, leading=13, alignment=align, textColor=colors.HexColor("#555555")
    )
    head = ParagraphStyle("h", fontName="Plex-Bold", fontSize=11, leading=15, alignment=align, spaceBefore=8)
    body = ParagraphStyle("b", fontName="Plex", fontSize=9, leading=13, alignment=align)
    cell = ParagraphStyle("c", fontName="Plex", fontSize=7.5, leading=10, alignment=align)
    cellb = ParagraphStyle("cb", parent=cell, fontName="Plex-Bold")

    def p(text, style):
        from html import escape

        return Paragraph(escape(_shape(text, rtl), quote=False), style)

    story = [p(doc["title"], title), p(doc.get("subtitle", ""), sub), Spacer(1, 4 * mm)]
    for block in doc["blocks"]:
        kind = block[0]
        if kind == "heading":
            story.append(p(block[1], head))
        elif kind == "text":
            story.append(p(block[1], body))
        elif kind == "table":
            header, rows = block[1], block[2]
            data = [[p(h, cellb) for h in header]] + [[p(_cell(v), cell) for v in r] for r in rows]
            if rtl:
                data = [list(reversed(r)) for r in data]
            width = pagesize[0] - 28 * mm
            table = Table(data, colWidths=[width / len(header)] * len(header), repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EDE9E3")),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BBBBBB")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("TOPPADDING", (0, 0), (-1, -1), 2),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                    ]
                )
            )
            story += [table, Spacer(1, 3 * mm)]
    stamp = timezone.localtime().strftime("%d/%m/%Y %H:%M")
    story.append(p(("أُنشئ في " if rtl else "Generated ") + stamp, sub))
    pdf.build(story)
    return buf.getvalue()


def to_xlsx(doc: dict, lang="en") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = (doc.get("sheet") or "Report")[:31]
    ws.sheet_view.rightToLeft = lang == "ar"
    ws.append([doc["title"]])
    ws["A1"].font = Font(bold=True, size=14)
    if doc.get("subtitle"):
        ws.append([doc["subtitle"]])
    ws.append([])
    widths = {}
    for block in doc["blocks"]:
        kind = block[0]
        if kind in ("heading", "text"):
            ws.append([block[1]])
            if kind == "heading":
                ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        elif kind == "table":
            header, rows = block[1], block[2]
            ws.append(list(header))
            for i in range(1, len(header) + 1):
                c = ws.cell(row=ws.max_row, column=i)
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="EDE9E3")
            for r in rows:
                ws.append([_xl(v) for v in r])
                for i, v in enumerate(r, start=1):
                    if isinstance(v, Decimal):
                        ws.cell(row=ws.max_row, column=i).number_format = "0.000"
            ws.append([])
            for i, h in enumerate(header, start=1):
                widths[i] = max(
                    widths.get(i, 8), min(40, len(str(h)) + 2), *(min(40, len(_cell(r[i - 1])) + 2) for r in rows)
                )
    for i, w in widths.items():
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=False)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xl(value):
    if isinstance(value, Decimal | int | float) and not isinstance(value, bool):
        return value
    if hasattr(value, "strftime"):
        return value
    return _cell(value)


def render(doc: dict, fmt: str, lang="en") -> tuple[bytes, str, str]:
    """Returns (content, content type, file extension)."""
    if fmt == "xlsx":
        return to_xlsx(doc, lang), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    return to_pdf(doc, lang), "application/pdf", "pdf"
